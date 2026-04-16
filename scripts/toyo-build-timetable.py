#!/usr/bin/env python3

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


DAY_ORDER = ["月", "火", "水", "木", "金", "土"]
PERIODS = ["1", "2", "3", "4", "5", "6", "7"]
DAY_COLORS = {
    "月": "FFF3BF",
    "火": "D3F9D8",
    "水": "D0EBFF",
    "木": "FFE8CC",
    "金": "E5DBFF",
    "土": "FFCCD5",
}
MODE_COLORS = {
    "対面": "FFF9DB",
    "非オ": "E3FAFC",
    "非同": "E7F5FF",
}


def normalize_digits(value: str) -> str:
    return value.translate(str.maketrans("０１２３４５６７８９", "0123456789"))


def make_workbook(data: dict, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "時間割"
    ws.freeze_panes = "B4"
    ws.sheet_view.showGridLines = False

    thin = Side(style="thin", color="8A8F98")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_fill = PatternFill("solid", fgColor="4C6EF5")
    subtitle_fill = PatternFill("solid", fgColor="EEF2FF")
    row_header_fill = PatternFill("solid", fgColor="FFF0F6")
    summary_fill = PatternFill("solid", fgColor="F1F3F5")

    academic_year = data["academicYear"]
    student_name = data["studentName"]
    student_number = data["studentNumber"]
    courses = data["courses"]
    total_credits = sum(course["credits"] or 0 for course in courses)

    ws.merge_cells("A1:G1")
    ws["A1"] = f"{academic_year}年度 履修時間割"
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:G2")
    ws["A2"] = f"{student_name} ({student_number})  /  小学校の時間割ふう"
    ws["A2"].font = Font(bold=True, color="1F2937", size=11)
    ws["A2"].fill = subtitle_fill
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    ws["A3"] = "時限"
    ws["A3"].font = Font(bold=True)
    ws["A3"].fill = summary_fill
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A3"].border = border

    for index, day in enumerate(DAY_ORDER, start=2):
        cell = ws.cell(row=3, column=index, value=day)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=DAY_COLORS[day])
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    ws.column_dimensions["A"].width = 9
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 24

    for row_index, period in enumerate(PERIODS, start=4):
        period_cell = ws.cell(row=row_index, column=1, value=f"{period}限")
        period_cell.font = Font(bold=True)
        period_cell.fill = row_header_fill
        period_cell.alignment = Alignment(horizontal="center", vertical="center")
        period_cell.border = border
        ws.row_dimensions[row_index].height = 68

        for day_index in range(2, 8):
            cell = ws.cell(row=row_index, column=day_index)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for course in courses:
        day = course["day"]
        period = normalize_digits(str(course["period"]))
        if day not in DAY_ORDER or period not in PERIODS:
            continue

        row = 4 + PERIODS.index(period)
        col = 2 + DAY_ORDER.index(day)
        room = course["room"] or "教室未設定"
        instructor = course["instructor"] or "担当者未設定"
        mode = course["deliveryMode"] or "未設定"
        credits = course["credits"] or ""
        cell = ws.cell(row=row, column=col)
        cell.value = (
            f"{course['courseName']}\n"
            f"{course['numbering']} / {course['courseCode']}\n"
            f"{instructor}\n"
            f"{room}・{course['campus']}\n"
            f"{mode} / {credits}単位"
        )
        cell.fill = PatternFill("solid", fgColor=MODE_COLORS.get(mode, "FFFFFF"))
        cell.font = Font(name="Meiryo", size=10, bold=True)

    summary_start = 12
    ws.merge_cells(start_row=summary_start, start_column=1, end_row=summary_start, end_column=7)
    ws.cell(row=summary_start, column=1, value="履修サマリー").fill = summary_fill
    ws.cell(row=summary_start, column=1).font = Font(bold=True)
    ws.cell(row=summary_start, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[summary_start].height = 22

    summary_lines = [
        f"登録科目数: {len(courses)}件",
        f"登録単位数: {total_credits}単位",
        f"実施形態: " + " / ".join(
            f"{mode} {sum(1 for course in courses if course['deliveryMode'] == mode)}件"
            for mode in sorted({course['deliveryMode'] for course in courses})
        ),
        "注: 学務ポータルの「履修登録確認表照会」をもとに作成",
    ]
    for offset, line in enumerate(summary_lines, start=1):
        ws.merge_cells(
            start_row=summary_start + offset,
            start_column=1,
            end_row=summary_start + offset,
            end_column=7,
        )
        summary_cell = ws.cell(row=summary_start + offset, column=1, value=line)
        summary_cell.alignment = Alignment(horizontal="left", vertical="center")
        summary_cell.border = border
        summary_cell.fill = summary_fill

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: toyo-build-timetable.py <input-json> <output-xlsx>", file=sys.stderr)
        return 1

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    data = json.loads(input_path.read_text(encoding="utf-8"))
    make_workbook(data, output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
